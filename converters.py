"""
تبدیل‌دهنده‌های اضافی ربات (غیر از Markdown → PDF که در bot.py است).

    md_to_docx(...)              Markdown  → Word (.docx)   با pandoc
    build_reference_docx(...)    ساخت قالب Word فارسی (RTL، فونت Vazirmatn)
    docx_to_pdf(...)             Word (.docx) → PDF        با LibreOffice (headless)

این ماژول هیچ وابستگی به bot.py ندارد. پیش‌پردازش مشترک Markdown (فرمول‌ها و ...)
را bot.py قبل از صدا زدن این توابع انجام می‌دهد.
"""
import os
import re
import shutil
import subprocess
import tempfile

# ───────────────────────────── یافتن باینری‌ها ─────────────────────────────


def find_pandoc():
    """ابتدا pandoc همراه pypandoc-binary، و اگر نبود pandoc سیستمی."""
    try:
        import pypandoc

        return pypandoc.get_pandoc_path()
    except Exception:
        found = shutil.which("pandoc")
        if not found:
            raise RuntimeError("pandoc پیدا نشد (pypandoc-binary نصب نیست)")
        return found


def find_soffice():
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    raise RuntimeError("LibreOffice (soffice) پیدا نشد")


# ───────────────────────────── Markdown → DOCX ─────────────────────────────

PANDOC_FROM = "markdown-raw_html+tex_math_dollars+pipe_tables+fenced_code_blocks+strikeout+footnotes"

CALLOUT_WORD_LABELS = {
    "NOTE": "📌 نکته:",
    "WARNING": "⚠️ هشدار:",
    "TIP": "💡 پیشنهاد:",
    "IMPORTANT": "🔥 مهم:",
    "EXAMPLE": "🩺 مثال بالینی:",
    "CLINICAL": "🧠 نکته بالینی:",
    "REVIEW": "❓ سؤال مرور:",
}
_CALLOUT_START_RE = re.compile(r"^\s*>\s*\[!(NOTE|WARNING|TIP|IMPORTANT|EXAMPLE|CLINICAL|REVIEW)\]\s*$", re.IGNORECASE)
_QUOTE_PREFIX_RE = re.compile(r"^\s*>\s?")


_NON_PLAIN_LINE_RE = re.compile(r"^(?:\s|[-*+]\s|\d+[.)]\s|>|\||#|<)")


def _split_soft_line_breaks(md_text: str) -> str:
    """
    سطرهای پشت‌سرهم یک پاراگراف (شکست نرم Markdown) را به پاراگراف جدا تبدیل می‌کند،
    تا هر خط جهت خودش را در Word بگیرد. فهرست، جدول، نقل‌قول، تیتر، کد و فرمول دست‌نخورده می‌مانند.
    """
    out = []
    in_fence = False
    in_math = False
    prev_plain = False
    for line in md_text.split("\n"):
        stripped = line.strip()
        if stripped.startswith(("```", "~~~")):
            in_fence = not in_fence
            out.append(line)
            prev_plain = False
            continue
        if in_fence:
            out.append(line)
            continue
        has_math_marker = "$$" in line
        if has_math_marker and line.count("$$") % 2 == 1:
            in_math = not in_math
        plain = (
            bool(stripped)
            and not in_math
            and not has_math_marker
            and not _NON_PLAIN_LINE_RE.match(line)
        )
        if plain and prev_plain:
            out.append("")
        out.append(line)
        prev_plain = plain
    return "\n".join(out)


def _docx_friendly_markdown(md_text: str) -> str:
    """
    callout های `> [!NOTE]` را به یک برچسب بولد و پاراگراف‌های عادی تبدیل می‌کند.
    (نقل‌قول‌های چندخطی بدون `>` اضافه می‌شوند تا در Word متن تمیز بماند.)
    """
    out = []
    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        match = _CALLOUT_START_RE.match(lines[i])
        if not match:
            out.append(lines[i])
            i += 1
            continue
        label = CALLOUT_WORD_LABELS[match.group(1).upper()]
        out.append(f"**{label}**")
        out.append("")
        i += 1
        while i < len(lines) and lines[i].lstrip().startswith(">"):
            out.append(_QUOTE_PREFIX_RE.sub("", lines[i], count=1))
            i += 1
        out.append("")
    return "\n".join(out)


def md_to_docx(md_text: str, out_path: str, title: str = "", reference_doc: str = None, timeout: int = 90) -> str:
    """Markdown را به .docx تبدیل می‌کند. فرمول‌ها به معادلهٔ قابل ویرایش Word تبدیل می‌شوند."""
    pandoc = find_pandoc()
    cmd = [
        pandoc,
        "--from", PANDOC_FROM,
        "--to", "docx",
        "--output", out_path,
        "--metadata", "lang=fa",
        "--metadata", "dir=rtl",
    ]
    if title:
        cmd += ["--metadata", f"title={title}"]
    if reference_doc and os.path.isfile(reference_doc):
        cmd += ["--reference-doc", reference_doc]

    source = _split_soft_line_breaks(_docx_friendly_markdown(md_text))
    result = subprocess.run(
        cmd,
        input=source.encode("utf-8"),
        capture_output=True,
        timeout=timeout,
    )
    if result.returncode != 0 or not os.path.isfile(out_path):
        raise RuntimeError(f"pandoc failed: {result.stderr.decode('utf-8', 'replace')[:300]}")
    apply_docx_directions(out_path)
    apply_docx_metadata(out_path, title)
    return out_path


# ───────────────────────────── قالب Word فارسی ─────────────────────────────

FARSI_FONT = "Vazirmatn"


def _insert_ordered(parent, element, successors):
    """عنصر را در جایگاه درست طبق ترتیب شِمای OOXML درج می‌کند (Word نسبت به ترتیب سخت‌گیر است)."""
    from docx.oxml.ns import qn

    for tag in successors:
        found = parent.find(qn(tag))
        if found is not None:
            found.addprevious(element)
            return element
    parent.append(element)
    return element


PPR_BIDI_SUCCESSORS = (
    "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
    "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)
RPR_RTL_SUCCESSORS = ("w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath")
SECT_BIDI_SUCCESSORS = ("w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange")
PERSIAN_CHAR_RE = re.compile(r"[\u0600-\u06FF\u0750-\u077F\uFB50-\uFDFF\uFE70-\uFEFF]")
LATIN_LETTER_RE = re.compile(r"[A-Za-z]")


def _docx_text_direction(text):
    """همان قاعده‌ای که در PDF اعمال می‌شود: فارسی در متن → راست‌به‌چپ، فقط لاتین → چپ‌به‌راست."""
    if PERSIAN_CHAR_RE.search(text or ""):
        return "rtl"
    if LATIN_LETTER_RE.search(text or ""):
        return "ltr"
    return "rtl"


def apply_docx_directions(docx_path: str) -> str:
    """برای هر پاراگراف جهت را بر اساس متنش تنظیم می‌کند (شامل خانه‌های جدول)."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document(docx_path)

    def fix(paragraph):
        ppr = paragraph._p.get_or_add_pPr()
        bidi = ppr.find(qn("w:bidi"))
        want_rtl = _docx_text_direction(paragraph.text) == "rtl"
        if want_rtl and bidi is None:
            _insert_ordered(ppr, OxmlElement("w:bidi"), PPR_BIDI_SUCCESSORS)
        elif not want_rtl and bidi is not None:
            ppr.remove(bidi)

    for paragraph in doc.paragraphs:
        fix(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    fix(paragraph)
    doc.save(docx_path)
    return docx_path


BRAND_NAME = "Medoxin"
CHANNEL_HANDLE = "@Med402"
CHANNEL_URL = "https://t.me/Med402"


def apply_docx_metadata(docx_path: str, title: str = "") -> str:
    """خواص فایل Word (Author، Subject، ...) را روی برند کانال می‌گذارد."""
    from docx import Document

    doc = Document(docx_path)
    props = doc.core_properties
    props.title = title or ""
    props.author = BRAND_NAME
    props.last_modified_by = BRAND_NAME
    props.subject = CHANNEL_HANDLE
    props.keywords = f"{BRAND_NAME}, {CHANNEL_HANDLE}"
    props.comments = CHANNEL_URL
    props.category = "جزوه"
    doc.save(docx_path)
    return docx_path


def build_reference_docx(out_path: str, font_name: str = FARSI_FONT) -> str:
    """
    قالب پیش‌فرض pandoc را می‌گیرد و آن را برای فارسی تنظیم می‌کند:
    همهٔ استایل‌های پاراگراف راست‌به‌چپ، فونت فارسی برای بخش‌های complex-script،
    و جهت راست‌به‌چپ برای بخش صفحه.
    """
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pandoc = find_pandoc()
    default_path = out_path + ".default.docx"
    with open(default_path, "wb") as fh:
        subprocess.run([pandoc, "--print-default-data-file", "reference.docx"], stdout=fh, check=True, timeout=60)

    doc = Document(default_path)

    def set_fonts(rpr):
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        # متن لاتین همیشه Arial؛ متن فارسی (cs/eastAsia) فونت فارسی
        for attr in ("w:ascii", "w:hAnsi"):
            rfonts.set(qn(attr), "Arial")
        for attr in ("w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), font_name)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]

    for style in doc.styles:
        try:
            element = style.element
            if style.type == WD_STYLE_TYPE.PARAGRAPH:
                ppr = element.get_or_add_pPr()
                if ppr.find(qn("w:bidi")) is None:
                    _insert_ordered(ppr, OxmlElement("w:bidi"), PPR_BIDI_SUCCESSORS)
            if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
                rpr = element.get_or_add_rPr()
                set_fonts(rpr)
                if rpr.find(qn("w:rtl")) is None:
                    _insert_ordered(rpr, OxmlElement("w:rtl"), RPR_RTL_SUCCESSORS)
        except Exception:
            continue

    for section in doc.sections:
        sectpr = section._sectPr
        if sectpr.find(qn("w:bidi")) is None:
            _insert_ordered(sectpr, OxmlElement("w:bidi"), SECT_BIDI_SUCCESSORS)

    doc.save(out_path)
    try:
        os.remove(default_path)
    except OSError:
        pass
    return out_path


# ───────────────────────────── فونت static برای LibreOffice ─────────────────────────────
# LibreOffice با فونت variable خوب کار نمی‌کند (وزن‌ها اعمال نمی‌شوند و بولد مصنوعی می‌شود).
# این تابع از فونت variable دو فایل static (Regular و Bold) می‌سازد.


def _set_static_names(font, family, style, bold):
    name = font["name"]
    for name_id in (1, 2, 3, 4, 6, 16, 17):
        name.removeNames(nameID=name_id)
    full = f"{family} {style}"
    ps_name = f"{family}-{style}"
    for platform, encoding, language in ((3, 1, 0x409), (1, 0, 0)):
        name.setName(family, 1, platform, encoding, language)
        name.setName(style, 2, platform, encoding, language)
        name.setName(f"{ps_name};static", 3, platform, encoding, language)
        name.setName(full, 4, platform, encoding, language)
        name.setName(ps_name, 6, platform, encoding, language)

    os2 = font["OS/2"]
    os2.usWeightClass = 700 if bold else 400
    selection = os2.fsSelection & ~(0x01 | 0x20 | 0x40)
    os2.fsSelection = selection | (0x20 if bold else 0x40)

    head = font["head"]
    head.macStyle = (head.macStyle | 0x01) if bold else (head.macStyle & ~0x01)

    if "STAT" in font:
        del font["STAT"]


def make_static_fonts(variable_ttf, out_dir, family="Vazirmatn"):
    """از فونت variable، فایل‌های static Regular و Bold می‌سازد و مسیرها را برمی‌گرداند."""
    from fontTools.ttLib import TTFont
    from fontTools.varLib import instancer

    os.makedirs(out_dir, exist_ok=True)
    result = {}
    for style, weight, bold in (("Regular", 400, False), ("Bold", 700, True)):
        font = TTFont(variable_ttf)
        if "fvar" in font:
            font = instancer.instantiateVariableFont(font, {"wght": weight}, updateFontNames=False)
        _set_static_names(font, family, style, bold)
        out_path = os.path.join(out_dir, f"{family}-{style}.ttf")
        font.save(out_path)
        result[style] = out_path
    return result


# ───────────────────────────── DOCX → PDF (LibreOffice) ─────────────────────────────
# این بخش از همین حالا آماده است؛ Dockerfile و requirements.txt با LibreOffice و
# فونت‌ها آماده‌اند و برای فعال کردن آن نیازی به تغییر آن‌ها نیست.


def docx_to_pdf(docx_path: str, out_dir: str, timeout: int = 120) -> str:
    """Word را با LibreOffice به PDF تبدیل می‌کند و مسیر PDF را برمی‌گرداند."""
    soffice = find_soffice()
    profile_dir = tempfile.mkdtemp(prefix="lo_profile_")
    try:
        cmd = [
            soffice,
            f"-env:UserInstallation=file://{profile_dir}",
            "--headless",
            "--norestore",
            "--convert-to", "pdf",
            "--outdir", out_dir,
            docx_path,
        ]
        result = subprocess.run(cmd, capture_output=True, timeout=timeout)
        base = os.path.splitext(os.path.basename(docx_path))[0]
        pdf_path = os.path.join(out_dir, base + ".pdf")
        if result.returncode != 0 or not os.path.isfile(pdf_path):
            raise RuntimeError(f"LibreOffice failed: {result.stderr.decode('utf-8', 'replace')[:300]}")
        return pdf_path
    finally:
        shutil.rmtree(profile_dir, ignore_errors=True)
